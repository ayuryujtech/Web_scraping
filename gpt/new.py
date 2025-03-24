import openai
import openpyxl
import time
import re
import json
from openpyxl import load_workbook

# Set your OpenAI API key
openai.api_key = ' '

# List of predefined categories
categories = [
    "ayurveda products", "vitamins & supplements", "pain relief", "multivitamins", "health & wellness", 
    "juices & vinegars", "women's health", "covid essentials", "ayurveda", "orthotics", "women care", 
    "women multivitamins", "hypertension products", "daily essentials", "personal care", "stomach care", 
    "calcium products", "omega", "adult diapers", "herbal supplements", "adult hygiene", "juices", 
    "health essentials", "ayurvedic medicine", "bone & joint health", "weight management", "elderly care", 
    "personal hygiene", "holistic wellness", "winter essentials", "masks", "herbal juices", 
    "bone & joint supplements", "hair & skin supplements", "heart supplements", "heart & bone health", 
    "immunity & wellness", "nutrition", "nutritional supplements", "food & nutrition", "cosmetics", 
    "skin care", "nutritional drinks", "fitness supplements", "feminine care", "healthy snacks", 
    "diabetes", "immunity boosters", "hair care", "sexual wellness", "eye care", "liver care", 
    "bone, joint & muscle care", "kidney care", "respiratory care", "homeopathy", "health care", 
    "malaria", "baby care", "oral care", "mindcare", "heart care", "derma care", "men grooming", 
    "pet care", "cardiac care", "cough & cold", "winter popular products", "headache & fever"
]

# Function to get predictions from OpenAI
def get_predictions(title, description):
    prompt = f"""
     Based on the product's title, description, and key ingredients, categorize the product into relevant categories from this list: {', '.join(categories)}.
    Provide exactly 5 key benefits with short and concise descriptions. All 5 benefits must be included and complete.

    Generate a list of icon ideas for Ayurveda products that emphasize the themes of nature, immunity, and personal well-being. For each benefit, provide:

    A brief description of the icon design.
    The meaning or significance of the icon in relation to Ayurveda.
    
    For each key benefit generated, also provide an icon name related to it based on these categories:
    
    Nature-Related Icons: Herbal Extracts, Forest-Sourced, Earth-Friendly, Pure Water, Sun-Powered.
    Immunity-Boosting Icons: Strength & Vitality, Immunity Shield, Energizing, Anti-Oxidant Rich, Detoxifying.
    Personal Well-being Icons: Holistic Health, Self-Care, Relaxation, Personalized Care, Natural Healing.
    "Nature-Related Icons: Earth-Friendly dont give like this give only Earth-Friendly "

    Title: {title}
    Description: {description}

    Respond in strict format as:
    {{
        "Categories": ["Category1", "Category2"],
        "Key Benefits": [
            {{
                "heading": "Benefit Heading 1",
                "description": "A short benefit description (10-20 words).",
                "cardType": "medium",
                "icon": "Generated icon name based on the benefit"
            }},
            {{
                "heading": "Benefit Heading 2",
                "description": "A short benefit description (10-20 words).",
                "cardType": "medium",
                "icon": "Generated icon name based on the benefit"
            }},
            {{
                "heading": "Benefit Heading 3",
                "description": "A short benefit description (10-20 words).",
                "cardType": "medium",
                "icon": "Generated icon name based on the benefit"
            }},
            {{
                "heading": "Benefit Heading 4",
                "description": "A short benefit description (10-20 words).",
                "cardType": "medium",
                "icon": "Generated icon name based on the benefit"
            }},
            {{
                "heading": "Benefit Heading 5",
                "description": "A longer benefit description (20-30 words).",
                "cardType": "large",
                "icon": "Generated icon name based on the benefit"
            }}
        ]
    }}
    """

    try:
        response = openai.ChatCompletion.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "You are a helpful assistant that classifies products and generates benefits."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=500,
            temperature=0.7
        )

        content = response['choices'][0]['message']['content'].strip()
        if not content:
            print("Received empty content. Skipping.")
            return None
        
        return content
    
    except openai.error.OpenAIError as e:
        print(f"OpenAI API error occurred: {e}")
        time.sleep(60)  # Wait before retrying
        return get_predictions(title, description)

# Function to extract categories and benefits from the GPT response
def extract_categories_and_benefits(text):
    try:
        categories_match = re.search(r'\"Categories\": \[([^\]]+)\]', text)
        categories = [cat.strip().strip('"') for cat in categories_match.group(1).split(',')] if categories_match else []

        benefits_match = re.findall(r'\{\s*"heading":\s*"([^"]+)",\s*"description":\s*"([^"]+)",\s*"cardType":\s*"([^"]+)",\s*"icon":\s*"([^"]+)"\s*\}', text)
        benefits = [
            {
                "heading": heading,
                "description": description,
                "cardType": cardType,
                "icon": icon
            }
            for heading, description, cardType, icon in benefits_match
        ]

        return categories, benefits

    except Exception as e:
        print(f"Error extracting categories and benefits: {e}")
        return [], []

# Function to process the Excel file and update with predictions
def process_excel(file_path):
    wb = load_workbook(file_path)
    sheet = wb.active
    
    sheet['J1'] = 'Predicted Categories'
    sheet['K1'] = 'Showcase Benefits'

    for row in range(2, sheet.max_row + 1):
        title = sheet[f'B{row}'].value  
        description = sheet[f'C{row}'].value  

        print(f"Processing row {row}: ")

        if not title or not description:
            print(f"Missing data in row {row}. Skipping.")
            continue

        predictions = get_predictions(title, description)
        
        if predictions:
            try:
                categories, benefits = extract_categories_and_benefits(predictions)

                if categories:
                    sheet[f'J{row}'] = ', '.join(categories)

                benefits_json = json.dumps(benefits, indent=4)
                sheet[f'K{row}'] = benefits_json
                
            except Exception as e:
                print(f"Error processing row {row}: {e}")
                continue  

    wb.save(file_path)
    print("All rows processed successfully and workbook saved.")

# Example usage
process_excel("/home/vedant/ScrapData/scrap/SriSriTatvaProduct.xlsx")
