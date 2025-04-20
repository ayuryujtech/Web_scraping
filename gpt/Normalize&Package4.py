import openpyxl
import json
import re
from openpyxl import load_workbook, Workbook

# Function to clean Unique Title by removing text after pipe symbol
def clean_unique_title(title):
    if isinstance(title, str) and '|' in title:
        return title.split('|')[0].strip()
    return title

# Function to normalize "Pack Size MRP" data
def normalize_pack_data(pack_data, base_product_name=""):
    normalized_data = []
    seen_variants = set()  # Track unique (size, flavor) tuples

    for idx, item in enumerate(pack_data):
        # Extract quantity and unit from size field using regex
        match = re.match(r"(\d+)\s*(\w+)", item.get('size', ''))
        if match:
            quantity, unit = match.groups()
            mrp = item.get('mrp', '')
            flavor = item.get('flavor', '')

            # Unique identifier for this variant (size and flavor)
            variant_key = (quantity, unit, flavor)
            if variant_key not in seen_variants:  # Check if variant is unique
                seen_variants.add(variant_key)  # Add to set of seen variants
                
                # Generate product name with quantity included
                product_name = f"{base_product_name} {quantity}{unit}"
                
                normalized_item = {
                    'type': 'primary' if idx == 0 else 'secondary',
                    'productName': product_name,
                    'size': {
                        'quantity': quantity,
                        'unit': unit
                    },
                    'mrp': mrp,
                    'sellingPrice': mrp,
                    'variant': flavor
                }
                normalized_data.append(normalized_item)
    return normalized_data

# Function to process the Excel file and include all data with the new normalized column
def process_excel(file_path):
    wb = load_workbook(file_path)
    sheet = wb.active

    # Identify column positions
    header_row = sheet[1]
    headers = {cell.value: cell.column for cell in header_row}
    
    # Find Pack Size MRP column
    if "Pack Size MRP" not in headers:
        print("Error: Could not find 'Pack Size MRP' column in the sheet.")
        return
    pack_size_mrp_col = headers["Pack Size MRP"]
    
    # Check if Unique Title column exists
    unique_title_col = headers.get("Unique Title")
    if unique_title_col:
        print("Found 'Unique Title' column - will clean text after '|' symbol")
    
    # Create a new workbook to save the processed data
    new_wb = Workbook()
    new_sheet = new_wb.active

    # Copy original headers and add new headers for normalized data
    header_values = [cell.value for cell in header_row]
    header_values.extend(["Normalized Data", "Variant Name"])
    new_sheet.append(header_values)

    # Process each row in the original Excel sheet
    for row in range(2, sheet.max_row + 1):
        row_data = [sheet.cell(row=row, column=col).value for col in range(1, sheet.max_column + 1)]
        
        # Clean Unique Title if column exists
        if unique_title_col:
            unique_title = sheet.cell(row=row, column=unique_title_col).value
            if unique_title:
                clean_title = clean_unique_title(unique_title)
                row_data[unique_title_col - 1] = clean_title
                if unique_title != clean_title:
                    print(f"Row {row}: Cleaned '{unique_title}' to '{clean_title}'")
        
        # Process Pack Size MRP data
        pack_size_mrp_raw = sheet.cell(row=row, column=pack_size_mrp_col).value

        if not pack_size_mrp_raw:
            row_data.extend(["", ""])  # Add empty normalized data and variant name if no data
        else:
            try:
                # Parse the JSON-like string in "Pack Size MRP" column
                pack_data = json.loads(pack_size_mrp_raw.replace("'", '"'))  # JSON parsing
                product_name = sheet.cell(row=row, column=headers["Product Name"]).value
                normalized_data = normalize_pack_data(pack_data, base_product_name=product_name)

                # Extract distinct flavors for the "Variant Name" column
                unique_flavors = {item.get('flavor', '') for item in pack_data}
                variant_name = ', '.join(filter(None, unique_flavors))  # Join unique flavors as a single string

                # Add normalized data and variant name to the row
                row_data.extend([json.dumps(normalized_data, indent=4), variant_name])

            except json.JSONDecodeError as e:
                print(f"Error decoding JSON in row {row}: {e}")
                row_data.extend(["", ""])  # Add empty normalized data and variant name in case of error

        # Add the row to the new sheet
        new_sheet.append(row_data)

    # Save the new workbook
    new_file_path = file_path.replace(".xlsx", "_normalized.xlsx")
    new_wb.save(new_file_path)
    print(f"All rows processed successfully and new workbook saved at {new_file_path}")


# Call the function with the input file path
if __name__ == "__main__":
    input_file = r'E:\AYURYUJ\Web_scraping\all_excelFiles\kottakal\all_kottakal_Filled.xlsx'
    process_excel(input_file)
